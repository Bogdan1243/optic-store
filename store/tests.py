import io
import zipfile
import openpyxl
from django.test import TestCase, override_settings
from store.models import Product, ProductImage
from store.excel_import import import_products_from_excel, generate_sample_xlsx


@override_settings(DEFAULT_FILE_STORAGE='django.core.files.storage.FileSystemStorage')
class ExcelImportTestCase(TestCase):

    def test_generate_sample_xlsx(self):
        content = generate_sample_xlsx()
        self.assertTrue(len(content) > 0)
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertIn('brand', headers)
        self.assertIn('price', headers)
        self.assertIn('color', headers)
        self.assertIn('size', headers)

    def test_import_products_without_images(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['brand', 'model_code', 'color', 'size', 'price', 'category', 'sex', 'description'])
        ws.append(['Gucci', 'GG0001', 'Gold', '54-18', 299.99, 'SUN', 'Women', 'Luxury sunglasses'])

        excel_io = io.BytesIO()
        wb.save(excel_io)
        excel_io.seek(0)

        res = import_products_from_excel(excel_io)
        self.assertEqual(res['created'], 1)
        self.assertEqual(res['updated'], 0)
        self.assertEqual(len(res['errors']), 0)

        product = Product.objects.get(brand='Gucci', color='Gold')
        self.assertEqual(float(product.price), 299.99)
        self.assertEqual(product.model_code, 'GG0001')
        self.assertEqual(product.sex, 'Women')

    def test_import_products_with_zip_images(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['brand', 'color', 'size', 'price', 'image', 'ar_image', 'gallery_images'])
        ws.append(['Prada', 'Black', '52-20', 350.00, 'prada_front.jpg', 'prada_ar.png', 'prada_side.jpg'])

        excel_io = io.BytesIO()
        wb.save(excel_io)
        excel_io.seek(0)

        zip_io = io.BytesIO()
        with zipfile.ZipFile(zip_io, 'w') as zf:
            zf.writestr('prada_front.jpg', b'fake_jpg_data')
            zf.writestr('prada_ar.png', b'fake_png_data')
            zf.writestr('prada_side.jpg', b'fake_side_data')
        zip_io.seek(0)

        res = import_products_from_excel(excel_io, zip_file_or_path=zip_io)
        self.assertEqual(res['created'], 1)

        product = Product.objects.get(brand='Prada', color='Black')
        self.assertTrue(bool(product.image))
        self.assertTrue(bool(product.ar_image))

        gallery_count = ProductImage.objects.filter(product=product).count()
        self.assertEqual(gallery_count, 1)

    def test_import_updates_existing_product(self):
        Product.objects.create(
            brand='Tom Ford',
            color='Havana',
            size='51-19',
            price=200.00,
            category='SUN'
        )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['brand', 'color', 'size', 'price', 'category'])
        ws.append(['Tom Ford', 'Havana', '51-19', 249.99, 'SUN'])

        excel_io = io.BytesIO()
        wb.save(excel_io)
        excel_io.seek(0)

        res = import_products_from_excel(excel_io)
        self.assertEqual(res['created'], 0)
        self.assertEqual(res['updated'], 1)

        product = Product.objects.get(brand='Tom Ford', color='Havana')
        self.assertEqual(float(product.price), 249.99)

    def test_export_and_reimport_products(self):
        from store.excel_import import export_products_to_xlsx

        p = Product.objects.create(
            brand='Armani Exchange',
            model_code='0AX1019',
            color='Negru',
            size='M',
            price=400.00,
            category='OPT',
            bridge_length='145',
            temple_length=None,
        )

        xlsx_bytes = export_products_to_xlsx()
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertIn('id', headers)
        self.assertIn('bridge_length', headers)
        self.assertIn('temple_length', headers)

        bridge_col = headers.index('bridge_length') + 1
        temple_col = headers.index('temple_length') + 1

        # Fix bridge_length and temple_length in the exported sheet
        ws.cell(row=2, column=bridge_col, value='18')
        ws.cell(row=2, column=temple_col, value='145')

        updated_io = io.BytesIO()
        wb.save(updated_io)
        updated_io.seek(0)

        res = import_products_from_excel(updated_io)
        self.assertEqual(res['created'], 0)
        self.assertEqual(res['updated'], 1)

        p.refresh_from_db()
        self.assertEqual(p.bridge_length, '18')
        self.assertEqual(p.temple_length, '145')


class SearchTestCase(TestCase):

    def setUp(self):
        self.p1 = Product.objects.create(
            brand='Ray-Ban',
            model_code='0RB2132',
            color='Negru Mat',
            size='M',
            price=990.00,
            category='SUN',
        )
        self.p2 = Product.objects.create(
            brand='Armani Exchange',
            model_code='0AX1019',
            color='Negru',
            size='M',
            price=400.00,
            category='OPT',
        )

    def test_search_by_brand_name(self):
        response = self.client.get('/shop/', {'q': 'Ray-Ban'})
        self.assertEqual(response.status_code, 200)
        products = list(response.context['products'])
        self.assertEqual(products, [self.p1])

    def test_search_by_brand_name_without_hyphen(self):
        response = self.client.get('/shop/', {'q': 'Ray Ban'})
        self.assertEqual(response.status_code, 200)
        products = list(response.context['products'])
        self.assertEqual(products, [self.p1])

    def test_search_by_product_code(self):
        response = self.client.get('/shop/', {'q': 'RB2132'})
        self.assertEqual(response.status_code, 200)
        products = list(response.context['products'])
        self.assertEqual(products, [self.p1])

        response2 = self.client.get('/shop/', {'q': '0AX1019'})
        self.assertEqual(response2.status_code, 200)
        products2 = list(response2.context['products'])
        self.assertEqual(products2, [self.p2])


class FavouritesTestCase(TestCase):

    def setUp(self):
        self.product = Product.objects.create(
            brand='Ray-Ban',
            model_code='0RB2132',
            color='Negru Mat',
            size='M',
            price=990.00,
            category='SUN',
        )

    def test_toggle_favourite_and_view_page(self):
        # Add to favourites
        res1 = self.client.post(f'/favourites/toggle/{self.product.id}/')
        self.assertEqual(res1.status_code, 200)
        data1 = res1.json()
        self.assertTrue(data1['is_favourite'])
        self.assertEqual(data1['count'], 1)

        # View favourites page
        fav_page = self.client.get('/favourites/')
        self.assertEqual(fav_page.status_code, 200)
        self.assertEqual(list(fav_page.context['products']), [self.product])

        # Remove from favourites
        res2 = self.client.post(f'/favourites/toggle/{self.product.id}/')
        self.assertEqual(res2.status_code, 200)
        data2 = res2.json()
        self.assertFalse(data2['is_favourite'])
        self.assertEqual(data2['count'], 0)


class ReturnPageAndFormTestCase(TestCase):

    def test_livrare_si_retur_page_renders(self):
        response = self.client.get('/livrare-si-retur/')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Termenul de retur')
        self.assertContains(response, 'Condiții de acceptare a returului')
        self.assertContains(response, 'Excepții de la dreptul de retur')
        self.assertContains(response, 'Formular de Retur')

    def test_formular_retur_page_renders(self):
        response = self.client.get('/formular-retur/')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Formular de Retur')

    def test_submit_valid_return_request(self):
        from store.models import ReturnRequest
        post_data = {
            'order_number': '#1055',
            'full_name': 'Ion Popescu',
            'email': 'ion@example.com',
            'phone': '0722123456',
            'products': '1x Ray-Ban RB2132 Negru Mat',
            'return_reason': 'nepotrivire',
            'refund_method': 'iban',
            'iban': 'RO49AAAA1B31007593840000',
            'account_holder': 'Ion Popescu',
            'observations': 'Coletul este pregătit.',
            'accept_conditions': 'on',
        }
        response = self.client.post('/livrare-si-retur/', post_data)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(ReturnRequest.objects.count(), 1)
        req = ReturnRequest.objects.first()
        self.assertEqual(req.order_number, '#1055')
        self.assertEqual(req.iban, 'RO49AAAA1B31007593840000')

    def test_iban_required_when_refund_method_is_iban(self):
        from store.models import ReturnRequest
        post_data = {
            'order_number': '#1056',
            'full_name': 'Ion Popescu',
            'email': 'ion@example.com',
            'phone': '0722123456',
            'products': '1x Ray-Ban RB2132',
            'return_reason': 'nepotrivire',
            'refund_method': 'iban',
            'iban': '',
            'accept_conditions': 'on',
        }
        response = self.client.post('/formular-retur/', post_data)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(ReturnRequest.objects.count(), 0)
        self.assertTrue(response.context['form'].errors.get('iban'))



