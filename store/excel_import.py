import os
import io
import zipfile
import openpyxl
from django.core.files.base import ContentFile
from . import models
from .models import Product

# Check if ProductImage model exists
HasProductImage = hasattr(models, 'ProductImage')
if HasProductImage:
    ProductImage = getattr(models, 'ProductImage')
else:
    ProductImage = None

# Header mapping to support English & Romanian column headers
HEADER_MAPPING = {
    'id': 'id',
    'brand': 'brand',
    'marca': 'brand',
    'model': 'model_code',
    'model_code': 'model_code',
    'cod_model': 'model_code',
    'cod model': 'model_code',
    'color': 'color',
    'culoare': 'color',
    'color_hex': 'color_hex',
    'cod_culoare': 'color_hex',
    'size': 'size',
    'marime': 'size',
    'mărime': 'size',
    'price': 'price',
    'pret': 'price',
    'preț': 'price',
    'stock': 'stock',
    'stoc': 'stock',
    'category': 'category',
    'categorie': 'category',
    'sex': 'sex',
    'gen': 'sex',
    'description': 'description',
    'descriere': 'description',
    'image': 'image',
    'main_image': 'image',
    'imagine': 'image',
    'imagine_principala': 'image',
    'ar_image': 'ar_image',
    'imagine_ar': 'ar_image',
    'ar': 'ar_image',
    'gallery_images': 'gallery_images',
    'imagini_galerie': 'gallery_images',
    'gallery': 'gallery_images',
    'frame_shape': 'frame_shape',
    'forma_cadru': 'frame_shape',
    'frame_material': 'frame_material',
    'material_cadru': 'frame_material',
    'frame_finish': 'frame_finish',
    'finisaj': 'frame_finish',
    'rim_type': 'rim_type',
    'tip_margine': 'rim_type',
    'bridge_length': 'bridge_length',
    'lungime_punte': 'bridge_length',
    'punte_nazala': 'bridge_length',
    'punte_nazală': 'bridge_length',
    'temple_length': 'temple_length',
    'lungime_brate': 'temple_length',
    'lungime_brațe': 'temple_length',
    'hinge_to_hinge': 'hinge_to_hinge',
    'balama_la_balama': 'hinge_to_hinge',
    'balama-la-balama': 'hinge_to_hinge',
    'lens_color': 'lens_color',
    'culoare_lentile': 'lens_color',
    'lens_width': 'lens_width',
    'latime_lentila': 'lens_width',
    'lățime_lentilă': 'lens_width',
    'latime_lentile': 'lens_width',
    'lățime_lentile': 'lens_width',
    'lens_height': 'lens_height',
    'inaltime_lentila': 'lens_height',
    'înălțime_lentilă': 'lens_height',
    'lens_base': 'lens_base',
    'baza_lentila': 'lens_base',
    'bază_lentilă': 'lens_base',
    'face_shape': 'face_shape',
    'forma_fetei': 'face_shape',
    'forma_feței': 'face_shape',
    'progressive_friendly': 'progressive_friendly',
    'reteta': 'progressive_friendly',
    'rețetă': 'progressive_friendly',
    'age_group': 'age_group',
    'grup_varsta': 'age_group',
    'grup_vârstă': 'age_group',
}


def normalize_header(raw_header):
    if not raw_header:
        return ''
    cleaned = str(raw_header).strip().lower().replace(' ', '_')
    return HEADER_MAPPING.get(cleaned, cleaned)


def generate_sample_xlsx():
    """Generates an in-memory sample Excel file template for product import."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Product Import Template"

    headers = [
        'brand', 'model_code', 'color', 'color_hex', 'size', 'price', 'stock',
        'category', 'sex', 'description', 'image', 'ar_image', 'gallery_images',
        'frame_shape', 'frame_material', 'frame_finish', 'rim_type',
        'bridge_length', 'temple_length', 'hinge_to_hinge', 'lens_color',
        'lens_width', 'lens_height', 'lens_base', 'face_shape',
        'progressive_friendly', 'age_group'
    ]

    ws.append(headers)

    # Example Row 1
    ws.append([
        'Ray-Ban', 'RB2140', 'Black', '#000000', '50-22', 150.00, 1,
        'SUN', 'Unisex', 'Classic Wayfarer sunglasses with green lenses.',
        'rb2140_black.jpg', 'rb2140_black_ar.png', 'rb2140_side.jpg, rb2140_angle.jpg',
        'Square', 'Acetate', 'Polished', 'Full Rim',
        '22', '150', '135', 'Green G-15',
        '50', '41', '6', 'Oval',
        True, 'Adult'
    ])

    # Example Row 2
    ws.append([
        'Oakley', 'OO9014', 'Matte Black', '#111111', '60-18', 180.00, 1,
        'OPT', 'Men', 'Durable optical frames designed for sport performance.',
        'oo9014_front.jpg', '', 'oo9014_side.jpg',
        'Rectangular', 'O Matter', 'Matte', 'Full Rim',
        '18', '138', '140', 'Clear',
        '60', '38', '4', 'Round',
        False, 'Adult'
    ])

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def export_products_to_xlsx(queryset=None):
    """Exports all products (or a provided queryset) into an Excel (.xlsx) file."""
    if queryset is None:
        queryset = Product.objects.all().order_by('id')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products Export"

    headers = [
        'id', 'brand', 'model_code', 'color', 'color_hex', 'size', 'price', 'stock',
        'category', 'sex', 'description', 'image', 'ar_image', 'gallery_images',
        'frame_shape', 'frame_material', 'frame_finish', 'rim_type',
        'bridge_length', 'temple_length', 'hinge_to_hinge', 'lens_color',
        'lens_width', 'lens_height', 'lens_base', 'face_shape',
        'progressive_friendly', 'age_group'
    ]
    ws.append(headers)

    for product in queryset:
        image_name = os.path.basename(product.image.name) if product.image else ''
        ar_image_name = os.path.basename(product.ar_image.name) if product.ar_image else ''
        gallery_names = ''
        if hasattr(product, 'images'):
            gallery_names = ', '.join(
                os.path.basename(gi.image.name)
                for gi in product.images.all()
                if gi.image
            )

        ws.append([
            product.id,
            product.brand or '',
            product.model_code or '',
            product.color or '',
            product.color_hex or '',
            product.size or '',
            float(product.price) if product.price is not None else 0.0,
            product.stock if product.stock is not None else 1,
            product.category or 'SUN',
            product.sex or 'Unisex',
            product.description or '',
            image_name,
            ar_image_name,
            gallery_names,
            product.frame_shape or '',
            product.frame_material or '',
            product.frame_finish or '',
            product.rim_type or '',
            product.bridge_length or '',
            product.temple_length or '',
            product.hinge_to_hinge or '',
            product.lens_color or '',
            product.lens_width or '',
            product.lens_height or '',
            product.lens_base or '',
            product.face_shape or '',
            bool(product.progressive_friendly),
            product.age_group or '',
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value if cell.value is not None else '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def _build_image_store(zip_file_or_path=None, image_dir=None):
    """
    Builds a dictionary mapping lowercased image filenames (and basenames)
    to a tuple of (filename, bytes_or_filepath).
    """
    image_store = {}

    if zip_file_or_path:
        # Safety limits to prevent ZIP bomb / memory exhaustion
        MAX_MEMBERS = 500
        MAX_FILE_SIZE = 20 * 1024 * 1024  # 20 MB per file
        MAX_TOTAL_SIZE = 500 * 1024 * 1024  # 500 MB total

        try:
            if isinstance(zip_file_or_path, (str, os.PathLike)):
                zf = zipfile.ZipFile(zip_file_or_path, 'r')
            else:
                zf = zipfile.ZipFile(zip_file_or_path, 'r')

            members = [m for m in zf.infolist() if not m.is_dir()]
            if len(members) > MAX_MEMBERS:
                raise ValueError(f"ZIP archive contains too many files ({len(members)} > {MAX_MEMBERS}).")

            total_size = sum(m.file_size for m in members)
            if total_size > MAX_TOTAL_SIZE:
                raise ValueError(f"ZIP archive uncompressed size too large ({total_size} bytes > {MAX_TOTAL_SIZE} bytes).")

            for member in members:
                if member.file_size > MAX_FILE_SIZE:
                    continue  # Skip oversized individual files

                filename = os.path.basename(member.filename)
                if not filename or filename.startswith('.'):
                    continue

                data = zf.read(member.filename)
                image_store[member.filename.lower()] = (filename, data)
                image_store[filename.lower()] = (filename, data)
        except ValueError:
            raise  # Re-raise our validation errors
        except Exception as e:
            raise ValueError(f"Could not process ZIP file: {e}")

    elif image_dir and os.path.exists(image_dir):
        for root, _, files in os.walk(image_dir):
            for file in files:
                if file.startswith('.'):
                    continue
                full_path = os.path.join(root, file)
                filename = os.path.basename(file)
                image_store[file.lower()] = (filename, full_path)
                image_store[filename.lower()] = (filename, full_path)

    return image_store


def _get_image_content(image_store, filename):
    """Retrieves (filename, ContentFile) for a given filename if present in image_store."""
    if not filename or not image_store:
        return None, None

    clean_name = str(filename).strip()
    key = clean_name.lower()

    if key not in image_store:
        basename_key = os.path.basename(clean_name).lower()
        if basename_key in image_store:
            key = basename_key
        else:
            return None, None

    target_name, data_or_path = image_store[key]

    if isinstance(data_or_path, bytes):
        return target_name, ContentFile(data_or_path, name=target_name)
    elif isinstance(data_or_path, str) and os.path.isfile(data_or_path):
        with open(data_or_path, 'rb') as f:
            return target_name, ContentFile(f.read(), name=target_name)

    return None, None


def import_products_from_excel(excel_file_or_path, zip_file_or_path=None, image_dir=None):
    """
    Parses an XLSX file and imports products into the database.
    """
    results = {
        'created': 0,
        'updated': 0,
        'errors': [],
        'warnings': [],
    }

    try:
        wb = openpyxl.load_workbook(excel_file_or_path, data_only=True)
    except Exception as e:
        results['errors'].append(f"Failed to open Excel file: {e}")
        return results

    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        results['errors'].append("Excel sheet is empty or contains no data rows.")
        return results

    raw_headers = rows[0]
    header_index_map = {}
    for idx, cell_val in enumerate(raw_headers):
        if cell_val is not None:
            norm_key = normalize_header(cell_val)
            header_index_map[norm_key] = idx

    required_columns = ['brand', 'color', 'size', 'price']
    missing_cols = [c for c in required_columns if c not in header_index_map]
    if missing_cols:
        results['errors'].append(f"Missing required columns in Excel: {', '.join(missing_cols)}")
        return results

    image_store = _build_image_store(zip_file_or_path, image_dir)

    CATEGORY_CHOICES_MAP = {
        'sun': 'SUN',
        'sunglasses': 'SUN',
        'opt': 'OPT',
        'optical': 'OPT',
        'optice': 'OPT',
        'ochelari de soare': 'SUN',
        'ochelari de vedere': 'OPT',
    }

    SEX_CHOICES_MAP = {
        'men': 'Men',
        'bărbați': 'Men',
        'bărbat': 'Men',
        'barbati': 'Men',
        'barbat': 'Men',
        'women': 'Women',
        'femei': 'Women',
        'femeie': 'Women',
        'children': 'Children',
        'copii': 'Children',
        'unisex': 'Unisex',
    }

    product_model_fields = {f.name: f for f in Product._meta.get_fields()}

    for row_idx, row_values in enumerate(rows[1:], start=2):
        if not any(row_values):
            continue

        def get_val(key, default=None):
            idx = header_index_map.get(key)
            if idx is not None and idx < len(row_values):
                val = row_values[idx]
                if val is not None:
                    return str(val).strip()
            return default

        brand = get_val('brand')
        color = get_val('color')
        size = get_val('size')
        price_raw = get_val('price')

        if not brand or not color or not size or price_raw is None:
            results['warnings'].append(f"Row {row_idx}: Skipped due to missing brand, color, size, or price.")
            continue

        try:
            price = float(price_raw)
        except ValueError:
            results['warnings'].append(f"Row {row_idx}: Invalid price value '{price_raw}'. Skipped row.")
            continue

        model_code = get_val('model_code', '')
        color_hex = get_val('color_hex', '')
        description = get_val('description', '')

        cat_raw = get_val('category', 'SUN').lower()
        category = CATEGORY_CHOICES_MAP.get(cat_raw, 'SUN')

        sex_raw = get_val('sex', 'Unisex').lower()
        sex = SEX_CHOICES_MAP.get(sex_raw, 'Unisex')

        prog_raw = get_val('progressive_friendly', 'false').lower()
        progressive_friendly = prog_raw in ['true', '1', 'yes', 'da', 't']

        stock_raw = get_val('stock')
        stock_val = None
        if stock_raw is not None and stock_raw != '':
            try:
                stock_val = max(0, int(float(stock_raw)))
            except ValueError:
                results['warnings'].append(f"Row {row_idx}: Invalid stock value '{stock_raw}', keeping default.")

        product = None
        id_raw = get_val('id')
        if id_raw:
            try:
                product = Product.objects.filter(pk=int(float(id_raw))).first()
            except ValueError:
                product = None

        if not product:
            product_filter = Product.objects.filter(
                brand__iexact=brand,
                color__iexact=color,
                size__iexact=size
            )
            if model_code and 'model_code' in product_model_fields:
                product_filter = product_filter.filter(model_code__iexact=model_code)

            product = product_filter.first()

        is_created = False

        if not product:
            product = Product(
                brand=brand,
                color=color,
                size=size,
                price=price,
                stock=stock_val if stock_val is not None else 1,
                description=description,
                category=category,
                sex=sex,
            )
            is_created = True
        else:
            product.brand = brand
            product.color = color
            product.size = size
            product.price = price
            if stock_val is not None:
                product.stock = stock_val
            if description or 'description' in header_index_map:
                product.description = description
            product.category = category
            product.sex = sex

        optional_fields = {
            'model_code': model_code,
            'color_hex': color_hex,
            'progressive_friendly': progressive_friendly,
            'frame_shape': get_val('frame_shape'),
            'frame_material': get_val('frame_material'),
            'frame_finish': get_val('frame_finish'),
            'rim_type': get_val('rim_type'),
            'bridge_length': get_val('bridge_length'),
            'temple_length': get_val('temple_length'),
            'hinge_to_hinge': get_val('hinge_to_hinge'),
            'lens_color': get_val('lens_color'),
            'lens_width': get_val('lens_width'),
            'lens_height': get_val('lens_height'),
            'lens_base': get_val('lens_base'),
            'face_shape': get_val('face_shape'),
            'age_group': get_val('age_group'),
        }

        for attr, val in optional_fields.items():
            if attr in product_model_fields:
                field_obj = product_model_fields[attr]
                if val is not None and val != '':
                    setattr(product, attr, val)
                elif attr in header_index_map and not is_created and getattr(field_obj, 'null', False):
                    setattr(product, attr, None)
                elif not getattr(field_obj, 'null', False):
                    if field_obj.get_internal_type() == 'BooleanField':
                        if getattr(product, attr, None) is None:
                            setattr(product, attr, False)

        # Handle Main Image
        image_name_raw = get_val('image')
        if image_name_raw and 'image' in product_model_fields:
            fname, content_file = _get_image_content(image_store, image_name_raw)
            if content_file:
                product.image.save(fname, content_file, save=False)
            elif image_store:
                results['warnings'].append(f"Row {row_idx}: Main image '{image_name_raw}' not found in provided image package.")

        # Handle AR Image
        ar_image_name_raw = get_val('ar_image')
        if ar_image_name_raw and 'ar_image' in product_model_fields:
            fname, content_file = _get_image_content(image_store, ar_image_name_raw)
            if content_file:
                product.ar_image.save(fname, content_file, save=False)
            elif image_store:
                results['warnings'].append(f"Row {row_idx}: AR image '{ar_image_name_raw}' not found in provided image package.")

        product.save()

        if is_created:
            results['created'] += 1
        else:
            results['updated'] += 1

        # Handle Gallery Images
        gallery_raw = get_val('gallery_images')
        if gallery_raw and ProductImage is not None:
            gallery_files = [f.strip() for f in gallery_raw.split(',') if f.strip()]
            for gfile in gallery_files:
                fname, content_file = _get_image_content(image_store, gfile)
                if content_file:
                    ProductImage.objects.create(product=product, image=content_file)
                elif image_store:
                    results['warnings'].append(f"Row {row_idx}: Gallery image '{gfile}' not found in image package.")

    return results
